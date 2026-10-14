from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sqlite3
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from ADECHIT.tools.preview import create_app
from routes.adechit.domain import DomainError
from routes.adechit.report_excel import MIMETYPE, xlsx_bytes


def workbook(model, columns=None):
    args = {} if columns is None else {'columns': columns}
    return load_workbook(BytesIO(xlsx_bytes(model, args))).active


@pytest.fixture
def model():
    return {'title': 'Report', 'filename': 'report', 'landscape': False, 'footer': 'Generated', 'blocks': [
        {'t': 'letterhead', 'left': ['Unit name', 'Full address'], 'right': 'October 2026'},
        {'t': 'title', 'text': 'Account', 'sub': ['01.10.2026 - 31.10.2026']},
        {'t': 'box', 'rows': [('CNP', '0123456789012'), ('Email', 'parent@example.test')]},
        {'t': 'table', 'cols': [{'label': 'Name', 'w': 3000}, {'label': 'CNP', 'w': 1800, 'align': 'c'},
                               {'label': 'Debit', 'w': 1000, 'align': 'r'}, {'label': 'Credit', 'w': 1000, 'align': 'r'}],
         'rows': [{'cells': ['=SUM(A1:A2)', '0123456789012', '1.234,50', '-50,25']},
                  {'cells': [{'t': 'TOTAL', 'span': 2}, '1.234,50', '-50,25'], 'cls': 'total'}]}]}


def values(sheet):
    return [cell.value for row in sheet for cell in row if cell.value is not None]


def test_default_keeps_all_columns_header_and_literal_identifiers(model):
    sheet = workbook(model)
    content = values(sheet)
    assert all(value in content for value in ['Unit name', 'Full address', 'October 2026', 'Account', 'CNP', 'parent@example.test'])
    row = next(row for row in sheet if row[0].value == '=SUM(A1:A2)')
    assert row[0].data_type == 's'
    assert row[1].value == '0123456789012' and row[1].data_type == 's'
    assert row[2].value == 1234.5 and row[2].data_type == 'n'
    assert row[3].value == -50.25
    assert model['blocks'][-1]['rows'][0]['cells'][2] == '1.234,50'


def test_column_filter_preserves_identity_and_projects_merged_totals(model):
    original = deepcopy(model)
    sheet = workbook(model, '0:1,0:3')
    content = values(sheet)
    assert 'Full address' in content and 'parent@example.test' in content
    assert 'Name' not in content and 'Debit' not in content and 1234.5 not in content
    total = next(row for row in sheet if row[0].value == 'TOTAL')
    assert total[1].value == -50.25
    assert model == original


def test_same_columns_apply_to_repeated_sections_and_headless_grand_total(model):
    table = deepcopy(model['blocks'][-1])
    model['blocks'].extend([{'t': 'break'}, {'t': 'letterhead', 'left': ['Second unit']}, table,
                            {**deepcopy(table), 'head': False, 'rows': [table['rows'][-1]]}])
    sheet = workbook(model, '0:0,0:2')
    content = values(sheet)
    assert content.count('Name') == 2 and content.count('Debit') == 2
    assert 'Credit' not in content and content.count('TOTAL') == 3
    assert content.count(1234.5) == 5 and 'Second unit' in content


@pytest.mark.parametrize('selection', ['', '0:99', '1:0', 'not-a-column'])
def test_invalid_or_empty_column_selection_is_rejected(model, selection):
    with pytest.raises(DomainError) as failure:
        workbook(model, selection)
    assert failure.value.status == 400 and failure.value.reason == 'EXPORT_COLUMNS'


def test_report_without_table_exports_complete_form(model):
    model['blocks'] = [block for block in model['blocks'] if block['t'] != 'table']
    model['blocks'].append({'t': 'sign', 'heads': ['Manager', 'Cashier']})
    assert 'Manager    |    Cashier' in values(workbook(model, ''))


@pytest.fixture
def http(tmp_path):
    # The existing preview initializer requires its September seed before building parent histories.
    app = create_app(tmp_path, closed_month=True)
    app.testing = True
    connection = sqlite3.connect(tmp_path / 'preview.sqlite')
    connection.executescript("""
        INSERT INTO AD_Plati (IDPL,IDP,IDZ,IDS,IDL,Data,Plata,TIP,Anulata) VALUES
            (20,1,1,1,3,'2026-10-10',125,1,0), (21,1,1,1,3,'2026-10-11',25,2,1);
        INSERT INTO AD_AlteDoc (IDPL,NrDoc,DataDoc,Explicatie,Anulata,IDL)
            VALUES (20,'OP-20','2026-10-10','Bank payment',0,3);
        INSERT INTO AD_Chitante (IDPL,Serie,Numar,Data,Anulata,IDL) VALUES (21,'DEMO',21,'2026-10-11',1,3);
    """)
    connection.commit()
    connection.close()
    return app.test_client()


@pytest.mark.parametrize('kind,params', [
    ('situation', 'month=3&group=1'), ('situation', 'month=3&all=1&split=1'),
    ('cash', 'start=2026-09-01&end=2026-10-31'), ('bank', 'start=2026-10-01&end=2026-10-31'),
    ('cancelled', 'month=3'), ('account', 'month=3&child=1'), ('account', 'month=3&all=1&group=1'),
    ('payment-order', 'id=1'),
])
def test_every_operator_report_offers_and_downloads_excel(http, kind, params):
    base = f'/api/adechit/reports/{kind}/'
    headers = {'X-Ade-Unit': 'preview'}
    preview = http.get(base + 'print?' + params, headers=headers)
    assert preview.status_code == 200, preview.json
    assert 'id="report-excel"' in preview.text and 'id="report-excel-dialog"' in preview.text
    response = http.get(base + 'xlsx?' + params, headers=headers)
    assert response.status_code == 200, response.json
    assert response.mimetype == MIMETYPE and '.xlsx' in response.headers['Content-Disposition']
    assert response.headers['Cache-Control'] == 'no-store'
    assert values(load_workbook(BytesIO(response.data)).active)


def test_export_filter_validation_and_read_permissions(http):
    url = '/api/adechit/reports/situation/xlsx?month=3&group=1'
    invalid = http.get(url + '&columns=', headers={'X-Ade-Unit': 'preview'})
    assert invalid.status_code == 400 and invalid.json['reason'] == 'EXPORT_COLUMNS'
    filtered = http.get(url + '&columns=0:1,0:6', headers={'X-Ade-Unit': 'preview'})
    assert filtered.status_code == 200
    content = values(load_workbook(BytesIO(filtered.data)).active)
    assert 'Unitate de probă' in content and 'Nume' in content and 'Contract\nstabilit' in content
    assert 'CNP' not in content and 'Total\nchitanțe' not in content
    http.application.config['ADE_TEST_RIGHTS'].discard('read')
    assert http.get(url, headers={'X-Ade-Unit': 'preview'}).status_code == 403


def test_parent_statement_keeps_masked_identity_with_one_selected_column():
    from routes.adechit.parent_statement import statement_model
    data = {'issuer': {'name': 'Unit', 'address': 'Address', 'tax_code': 'RO123', 'accounts': [{'Cont': 'RO001'}]},
            'child': 'Child', 'address': 'Parent address', 'cnp': '•••••••••1234', 'email': 'parent@example.test',
            'period': 'All', 'generated': '12.10.2026', 'rows': [
                {'date': '12.10.2026', 'explanation': 'Opening', 'debit': '1 234,50', 'credit': '0,00',
                 'balance': '1 234,50', 'balance_type': 'Debit', 'kind': 'total'}]}
    sheet = workbook(statement_model(data), '0:2')
    content = values(sheet)
    assert all(value in content for value in ['Child', 'Parent address', '•••••••••1234', 'Cont: RO001', 'Cod fiscal: RO123'])
    assert 'DEBIT' in content and 'CREDIT' not in content and 1234.5 in content
