"""PDF and simple XLSX render the exact same report model; no templates or SQL."""
from html import escape
from io import BytesIO
from pathlib import Path
from threading import Lock
import re
from decimal import Decimal
from reporting.excel_widths import ColumnWidths, column_key
from .errors import ReportError

FONT = Path(__file__).resolve().parents[1] / 'static' / 'fonts' / 'DejaVuSans.ttf'
FONT_LOCK = Lock()
GROUP_BACKGROUND = '#f7f8f2'
GRAND_TOTAL_GAP_PT = 2 * 72 / 25.4  # Same 2 mm separation as the category band.


def tokens(text, values):
    for key, value in values.items():
        text = str(text).replace('{' + key + '}', str(value or ''))
    return text


def pdf(model):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, Flowable, TopPadder
    from reportlab.pdfgen.canvas import Canvas
    with FONT_LOCK:
        if 'SharedReport' not in pdfmetrics.getRegisteredFontNames():
            if not FONT.is_file():
                raise ReportError('REPORT_RESOURCE_MISSING',
                               'Resursele raportului nu sunt instalate complet. '
                               'Lipsește static/fonts/DejaVuSans.ttf. '
                               'Contactați administratorul pentru republicarea resurselor.', 503)
            pdfmetrics.registerFont(TTFont('SharedReport', str(FONT)))
            pdfmetrics.registerFontFamily('SharedReport', normal='SharedReport', bold='SharedReport',
                                          italic='SharedReport', boldItalic='SharedReport')
    layout, identity = model['layout'], model['identity']
    page = landscape(A4) if layout['landscape'] else A4
    margin, size = 9 * mm, layout['font_size']
    width = page[0] - 2 * margin
    class EmphasizedParagraph(Paragraph):
        # Synthetic emphasis preserves the bundled Unicode font and diacritics.
        def draw(self):
            super().draw()
            self.canv.saveState()
            self.canv.translate(.22, 0)
            super().draw()
            self.canv.restoreState()
    def para(value, align=TA_LEFT, point=None, bold=False):
        style = ParagraphStyle('vc', fontName='SharedReport', fontSize=point or size,
                               leading=(point or size) * 1.22, alignment=align, spaceAfter=0)
        paragraph = EmphasizedParagraph if bold else Paragraph
        return paragraph(escape(str(value if value is not None else '')).replace('\n', '<br/>'), style)
    class PositionedBand(Flowable):
        def __init__(self, elements, available):
            super().__init__()
            self.width = available
            self.parts = []
            self.height = 0
            for element in elements:
                paragraph = para(element['text'], point=element.get('size', 8), bold=element.get('bold', False))
                _, height = paragraph.wrap(available * element['w'], 1000)
                self.parts.append((element, paragraph, height))
                self.height = max(self.height, element.get('y', 0) + max(element.get('h', 12), height))
        def draw(self):
            for element, paragraph, height in self.parts:
                paragraph.drawOn(self.canv, self.width * element['x'], self.height - element.get('y', 0) - height)
    def render(block, available):
        kind = block['t']
        if kind == 'form_grid':
            from . import form_grid
            return form_grid.pdf(block, available, para)
        if kind == 'positioned':
            return PositionedBand(block['elements'], available)
        if kind == 'break':
            return PageBreak()
        if kind == 'spacer':
            return Spacer(1, block.get('height', 5 * mm))
        if kind in ('title', 'text', 'group'):
            if kind == 'group':
                item = Table([[para(block['text'], TA_CENTER, size + .5, True)]], colWidths=[available])
                item.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), .4, colors.black), ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(GROUP_BACKGROUND))]))
                return item
            return para(block['text'], TA_CENTER if kind == 'title' else
                        {'c': TA_CENTER, 'r': TA_RIGHT}.get(block.get('align'), TA_LEFT),
                        layout['title_size'] if kind == 'title' else size, kind == 'title' or block.get('bold', False))
        if kind == 'box':
            data = [[para(label), para(value)] for label, value in block['rows']]
            item = Table(data, colWidths=[available * .34, available * .66])
        elif kind == 'sign':
            count = len(block['heads'])
            separated = count > 1 and block.get('box_width_ratio', 1) < 1
            if separated:
                box_width = available / count * block['box_width_ratio']
                gap_width = (available - box_width * count) / (count - 1)
                cells, widths = [], []
                for index, head in enumerate(block['heads']):
                    if index:
                        cells.append('')
                        widths.append(gap_width)
                    cells.append(para(head, TA_CENTER))
                    widths.append(box_width)
                item = Table([cells, [''] * len(cells)], colWidths=widths,
                             rowHeights=[None, 14 * mm], splitByRow=0)
                commands = []
                for index in range(0, len(cells), 2):
                    commands.extend([('BOX', (index, 0), (index, -1), .35, colors.black),
                                     ('LINEBELOW', (index, 0), (index, 0), .35, colors.black)])
                item.setStyle(TableStyle(commands))
            else:
                item = Table([[para(head, TA_CENTER) for head in block['heads']], [''] * count],
                             colWidths=[available / count] * count, rowHeights=[None, 14 * mm],
                             splitByRow=0 if block.get('anchor') == 'bottom' else 1)
        elif kind == 'columns':
            count = len(block['children'])
            # Flatten parallel tables into one splittable table. Nesting complete
            # tables in a single outer row would fail on a long list of foods.
            children = [next(part for part in child if part['t'] == 'table') for child in block['children']]
            columns, groups, offset = [], [], 0
            for child, parts in zip(children, block['children']):
                total = sum(c['w'] for c in child['cols'])
                columns.extend([{**c, 'w': c['w'] / total} for c in child['cols']])
                label = next((p['text'] for p in parts if p['t'] == 'group'), '')
                groups.append({'label': label, 'start': offset, 'span': len(child['cols'])})
                offset += len(child['cols'])
            rows = []
            for index in range(max(len(child['rows']) for child in children)):
                cells = []
                for child in children:
                    cells.extend(child['rows'][index]['cells'] if index < len(child['rows']) else [''] * len(child['cols']))
                rows.append({'cells': cells})
            return render({'t': 'table', 'cols': columns, 'rows': rows,
                           'show_column_header': any(child.get('show_column_header', True) for child in children),
                           'footer': block.get('footer'),
                           'section_boundaries': [group['start'] for group in groups[1:]],
                           'section_separator_width': block.get('section_separator_width'),
                           'grid_width': block.get('grid_width', .35),
                           'headers': [groups] if any(group['label'] for group in groups) else []}, available)
        elif kind == 'table':
            cols, commands, data, gap_rows = block['cols'], [], [], []
            headers = block.get('headers', [])
            for groups in headers:
                cells = [''] * len(cols)
                for group in groups:
                    start, span = group['start'], group['span']
                    cells[start] = para(group['label'], TA_CENTER)
                    if span > 1:
                        commands.append(('SPAN', (start, len(data)), (start + span - 1, len(data))))
                data.append(cells)
            if block.get('show_column_header', True):
                data.append([para(c['label'], TA_CENTER, c.get('header_size', c.get('size')), True) for c in cols])
            if block.get('formula'):
                data.append([para(value, TA_CENTER) for value in block['formula']])
                for index, value in enumerate(block['formula']):
                    if value == '' and len(data) == 2:
                        commands.append(('SPAN', (index, 0), (index, 1)))
            head_count = len(data)
            for row in block['rows']:
                if row.get('cls') == 'grand' and len(data) > head_count:
                    gap_rows.append(len(data))
                    data.append([''] * len(cols))
                    commands.append(('NOSPLIT', (0, len(data) - 2), (-1, len(data))))
                cells, index = [], 0
                for cell in row['cells']:
                    value, span = (cell['t'], cell.get('span', 1)) if isinstance(cell, dict) else (cell, 1)
                    col = cols[index]
                    cells.append(para(value, {'r': TA_RIGHT, 'c': TA_CENTER}.get(col.get('align'), TA_LEFT),
                                      col.get('size'), col.get('bold', False) or row.get('cls') in ('total', 'grand')
                                      or isinstance(cell, dict) and cell.get('bold', False)))
                    if isinstance(cell, dict) and cell.get('background'):
                        commands.append(('BACKGROUND', (index, len(data)), (index + span - 1, len(data)),
                                         colors.HexColor(cell['background'])))
                    if span > 1:
                        commands.append(('SPAN', (index, len(data)), (index + span - 1, len(data))))
                        cells.extend([''] * (span - 1))
                    index += span
                data.append(cells)
                if row.get('cls') == 'cancelled':
                    commands.append(('TEXTCOLOR', (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor('#7f7f7f')))
                if row.get('cls') in ('total', 'grand'):
                    background = GROUP_BACKGROUND if row.get('cls') == 'grand' else (
                        '#ffffff' if block.get('payment_notice') else '#ececec')
                    commands.append(('BACKGROUND', (0, len(data) - 1), (-1, len(data) - 1),
                                     colors.HexColor(background)))
                    # Keep each total with the preceding body row when paginating.
                    commands.append(('NOSPLIT', (0, max(head_count, len(data) - 2)), (-1, len(data) - 1)))
            footer = block.get('footer')
            if footer:
                data.append([para(footer['label'], TA_LEFT, footer.get('size'), True)] + [''] * (len(cols) - 2)
                            + [para(footer['value'], TA_RIGHT, footer.get('size'), True)])
                commands.extend([('SPAN', (0, len(data) - 1), (len(cols) - 2, len(data) - 1)),
                                 ('BACKGROUND', (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor('#ececec')),
                                 ('NOSPLIT', (0, max(0, len(data) - 2)), (-1, len(data) - 1))])
            if not data:
                # A headerless subreport may have no dishes at all.
                return Spacer(1, 0)
            denominator = sum(c['w'] for c in cols)
            actual_width = available * block.get('width', 1)
            item = Table(data, colWidths=[actual_width * c['w'] / denominator for c in cols],
                         rowHeights=[GRAND_TOTAL_GAP_PT if index in gap_rows else None for index in range(len(data))],
                         repeatRows=head_count, hAlign='CENTER')
            if head_count:
                commands.append(('BACKGROUND', (0, 0), (-1, head_count - 1),
                                 colors.white if block.get('payment_notice') else colors.HexColor('#ececec')))
                for index, column in enumerate(cols):
                    if column.get('header_background'):
                        commands.append(('BACKGROUND', (index, 0), (index, head_count - 1),
                                         colors.HexColor(column['header_background'])))
            if block.get('formula'):
                for index, value in enumerate(block['formula']):
                    if value:
                        commands.append(('BACKGROUND', (index, head_count - 1), (index, head_count - 1), colors.yellow))
            if block.get('section_separator_width'):
                commands.extend(('LINEBEFORE', (index, 0), (index, -1),
                                 block['section_separator_width'], colors.black)
                                for index in block.get('section_boundaries', []))
            item.setStyle(TableStyle(commands))
        else:
            raise ValueError('Unknown report band: ' + kind)
        grid = [] if kind == 'sign' and separated else [('GRID', (0, 0), (-1, -1), block.get('grid_width', .35), colors.black)]
        if kind == 'table' and gap_rows:
            grid, first = [], 0
            for boundary in gap_rows + [len(data)]:
                if boundary > first:
                    grid.append(('GRID', (0, first), (-1, boundary - 1), block.get('grid_width', .35), colors.black))
                first = boundary + 1
        item.setStyle(TableStyle(grid + [('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]))
        return TopPadder(item) if block.get('anchor') == 'bottom' else item
    elements = identity['elements'] + layout.get('elements', [])
    values = {'UNITATE': identity['name'], 'CUI': identity['fiscal_code'], 'ADRESA': identity['address'],
              'ANTET': identity['header'], 'SUBSOL': identity['footer'], 'GESTIUNE': model['gestiune_name'],
              'PERIOADA': model['period_text'], 'RAPORT': model['title'], 'SUBUNITATE': identity.get('subunit_name', '')}
    def band_height(band):
        ends = []
        for element in elements:
            if element['band'] != band:
                continue
            p = para(tokens(tokens(element['text'], {**values, 'PAGINA': 999, 'PAGINI': 999}), values),
                     point=element.get('size', 8), bold=element.get('bold', False))
            _, height = p.wrap(width * element['w'], 1000)
            ends.append(element.get('y', 0) + max(element.get('h', 12), height))
        return max(ends or [12])
    header_height, footer_height = band_height('header'), band_height('footer')
    class NumberedCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.states = []
        def showPage(self):
            self.states.append(dict(self.__dict__))
            self._startPage()
        def save(self):
            count = len(self.states)
            for state in self.states:
                self.__dict__.update(state)
                current = {**values, 'PAGINA': self._pageNumber, 'PAGINI': count}
                for e in elements:
                    if e['band'] not in ('header', 'footer'):
                        continue
                    y = page[1] - margin - float(e.get('y', 0)) if e['band'] == 'header' else margin + footer_height - float(e.get('y', 0))
                    p = para(tokens(tokens(e['text'], current), current),
                             {'r': TA_RIGHT, 'c': TA_CENTER}.get(e.get('align'), TA_LEFT),
                             point=e.get('size', 8), bold=e.get('bold', False))
                    w, h = p.wrap(width * e['w'], 100 * mm)
                    p.drawOn(self, margin + width * e['x'], y - h)
                Canvas.showPage(self)
            Canvas.save(self)
    story = []
    for block in model['blocks']:
        if block['t'] == 'richtext':
            from .richtext import pdf_flowables
            story.extend(pdf_flowables(block['paragraphs']))
            continue
        item = render(block, width)
        item.keepWithNext = block.get('keep_with_next', False)
        story.append(item)
        if block.get('anchor') != 'bottom' and block.get('space_after', 2):
            story.append(Spacer(1, block.get('space_after', 2) * mm))
    stream = BytesIO()
    SimpleDocTemplate(stream, pagesize=page, leftMargin=margin, rightMargin=margin,
        topMargin=margin + header_height + 6, bottomMargin=margin + footer_height + 6,
        title=model['title']).build(story, canvasmaker=NumberedCanvas)
    return stream.getvalue()


def xlsx(model, column_widths=None):
    from openpyxl import Workbook
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, column_index_from_string
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Raport'
    row_index = 1
    centered_rows = []
    full_rows = []
    form_headers = []
    widths = ColumnWidths(column_widths)
    edge = Side(style='thin', color='AAAAAA')
    def write(row, column, value, bold=False, point=10, numeric=False):
        cell = sheet.cell(row, column, str(value if value is not None else ''))
        # Text and identifiers stay literal. Only explicitly numeric columns
        # opt into the application's established Excel number conversion.
        cell.data_type = 's'
        if numeric and re.fullmatch(r'-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?', cell.value):
            text = cell.value
            cell.value = Decimal(text.replace('.', '').replace(',', '.'))
            cell.number_format = '#,##0.00' if ',' in text else '#,##0'
        cell.font = Font(name='Arial', size=point, bold=bold)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        return cell
    def table_footer(footer, count):
        nonlocal row_index
        write(row_index, 1, footer['label'], True, footer.get('size', 10))
        write(row_index, count, footer['value'], True, footer.get('size', 10))
        if count > 2:
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=count - 1)
        for column in range(1, count + 1):
            cell = sheet.cell(row_index, column)
            cell.fill = PatternFill('solid', fgColor='ECECEC')
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.alignment = Alignment(horizontal='right' if column == count else 'left', vertical='top')
        row_index += 1
    def blocks(items):
        nonlocal row_index
        for block in items:
            kind = block['t']
            if kind == 'form_grid':
                from . import form_grid
                form_headers.append((block, row_index))
                row_index = form_grid.xlsx(sheet, block, row_index, write, widths)
            elif kind == 'positioned':
                for element in sorted(block['elements'], key=lambda e: (e.get('y', 0), e.get('x', 0))):
                    write(row_index, 1, element['text'], element.get('bold', False), element.get('size', 10))
                    if block.get('full_width'):
                        full_rows.append((row_index, 'right' if element.get('x', 0) >= .6 else 'left'))
                    row_index += 1
            elif kind == 'break':
                sheet.row_breaks.append(Break(id=row_index - 1))
            elif kind == 'columns':
                for child in block['children']:
                    blocks(child)
                if block.get('footer'):
                    # Child tables leave two blank rows; the footer touches the last table.
                    row_index -= 2
                    count = max(len(part['cols']) for child in block['children'] for part in child if part['t'] == 'table')
                    table_footer(block['footer'], count)
                    row_index += 2
            elif kind in ('title', 'text', 'group'):
                write(row_index, 1, block['text'], kind != 'text' or block.get('bold', False))
                if block.get('full_width'):
                    full_rows.append((row_index, {'c': 'center', 'r': 'right'}.get(block.get('align'), 'left')))
                elif block.get('align') == 'c':
                    centered_rows.append(row_index)
                row_index += 1 if block.get('space_after') == 0 else 2
            elif kind == 'richtext':
                from .richtext import plain_text
                for paragraph in block['paragraphs']:
                    write(row_index, 1, plain_text([paragraph]))
                    full_rows.append((row_index, paragraph['align']))
                    row_index += 1
            elif kind in ('box', 'sign'):
                for cells in block.get('rows', [block.get('heads', [])]):
                    for i, value in enumerate(cells, 1):
                        write(row_index, i, value)
                    row_index += 1
                row_index += 1
            elif kind == 'table':
                for i, col in enumerate(block['cols'], 1):
                    if block.get('show_column_header', True):
                        cell = write(row_index, i, col['label'], True, col.get('header_size', col.get('size', 10)))
                        if 'header_size' in col:
                            cell.fill = PatternFill('solid', fgColor='ECECEC')
                        if col.get('header_background'):
                            cell.fill = PatternFill('solid', fgColor=col['header_background'].lstrip('#'))
                    dimension = sheet.column_dimensions[get_column_letter(i)]
                    dimension.width = widths.observe(i, column_key(col, i - 1), col['w'], dimension.width)
                if block.get('show_column_header', True):
                    row_index += 1
                if block.get('formula'):
                    for i, value in enumerate(block['formula'], 1):
                        write(row_index, i, value)
                    row_index += 1
                for record_index, record in enumerate(block['rows']):
                    if record.get('cls') == 'grand' and record_index:
                        sheet.row_dimensions[row_index].height = GRAND_TOTAL_GAP_PT
                        row_index += 1
                    i = 1
                    for value in record['cells']:
                        col = block['cols'][i - 1]
                        span = value.get('span', 1) if isinstance(value, dict) else 1
                        cell = write(row_index, i, value['t'] if isinstance(value, dict) else value,
                                     record.get('cls') in ('total', 'grand') or col.get('bold', False), col.get('size', 10),
                                     numeric=col.get('numeric', False))
                        cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                        cell.alignment = Alignment(horizontal={'r': 'right', 'c': 'center'}.get(col.get('align'), 'left'),
                                                   vertical='top', wrap_text=True)
                        if record.get('cls') == 'grand':
                            cell.fill = PatternFill('solid', fgColor=GROUP_BACKGROUND.lstrip('#'))
                        elif record.get('cls') == 'total':
                            cell.fill = PatternFill('solid', fgColor='ECECEC')
                        elif isinstance(value, dict) and value.get('background'):
                            cell.fill = PatternFill('solid', fgColor=value['background'].lstrip('#'))
                        if span > 1:
                            sheet.merge_cells(start_row=row_index, start_column=i,
                                              end_row=row_index, end_column=i + span - 1)
                        i += span
                    row_index += 1
                if block.get('footer'):
                    table_footer(block['footer'], len(block['cols']))
                row_index += 2
    identity = model['identity']
    values = {'UNITATE': identity['name'], 'CUI': identity['fiscal_code'], 'ADRESA': identity['address'],
              'GESTIUNE': model['gestiune_name'], 'PERIOADA': model['period_text'], 'RAPORT': model['title'],
              'ANTET': identity['header'], 'SUBSOL': identity['footer'], 'SUBUNITATE': identity.get('subunit_name', '')}
    def page_band(name):
        nonlocal row_index
        elements = [e for e in identity['elements'] + model['layout'].get('elements', []) if e['band'] == name]
        for element in sorted(elements, key=lambda e: (e.get('y', 0), e.get('x', 0))):
            text = tokens(tokens(element['text'], values), values)
            if '{PAGINA}' in text or '{PAGINI}' in text:
                part = sheet.oddHeader if name == 'header' else sheet.oddFooter
                position = part.right if element.get('x', 0) > .65 else part.left
                position.text = text.replace('{PAGINA}', '&P').replace('{PAGINI}', '&N')
            else:
                for line in text.splitlines() or ['']:
                    write(row_index, 1, line, element.get('bold', name == 'header'), element.get('size', 10))
                    if element.get('full_width'):
                        full_rows.append((row_index, 'right' if element.get('align') == 'r' else 'left'))
                    row_index += 1
        row_index += 1
    page_band('header')
    header_end = row_index - 1
    if identity.get('letterhead'):
        sheet.print_title_rows = f'1:{header_end}'
    blocks(model['blocks'])
    page_band('footer')
    for row in centered_rows:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=sheet.max_column)
        sheet.cell(row, 1).alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
    for row, align in full_rows:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=sheet.max_column)
        sheet.cell(row, 1).alignment = Alignment(horizontal=align, vertical='top', wrap_text=True)
    for dimension in sheet.column_dimensions.values():
        index = column_index_from_string(dimension.index)
        dimension.width = widths.final(index, dimension.width)
    if form_headers:
        from . import form_grid
        for block, first_row in form_headers:
            form_grid.fit_rows(sheet, block, first_row)
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = 'landscape' if model['layout']['landscape'] else 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.freeze_panes = f'A{header_end + 1}' if identity.get('letterhead') else 'A4'
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
