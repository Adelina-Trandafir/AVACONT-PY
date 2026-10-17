"""PDF and simple XLSX render the exact same report model; no templates or SQL."""
from html import escape
from io import BytesIO
from pathlib import Path
from threading import Lock
from reporting.excel_widths import ColumnWidths
from .nomenclatoare import NomError

FONT = Path(__file__).resolve().parents[2] / 'static' / 'fonts' / 'DejaVuSans.ttf'
FONT_LOCK = Lock()


def tokens(text, values):
    for key, value in values.items():
        text = str(text).replace('{' + key + '}', str(value or ''))
    return text


def pdf(model):
    if model['layout'].get('form') == 'nir-classic':
        from . import report_reception
        return report_reception.pdf(model)
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, Flowable
    from reportlab.pdfgen.canvas import Canvas
    with FONT_LOCK:
        if 'VerconReport' not in pdfmetrics.getRegisteredFontNames():
            if not FONT.is_file():
                raise NomError('REPORT_RESOURCE_MISSING',
                               'Resursele raportului nu sunt instalate complet. '
                               'Lipsește static/fonts/DejaVuSans.ttf. '
                               'Contactați administratorul pentru republicarea resurselor.', 503)
            pdfmetrics.registerFont(TTFont('VerconReport', str(FONT)))
    layout, identity = model['layout'], model['identity']
    page = landscape(A4) if layout['landscape'] else A4
    margin, size = 9 * mm, layout['font_size']
    width = page[0] - 2 * margin
    class EmphasizedParagraph(Paragraph):
        # Synthetic emphasis preserves the bundled Unicode font and diacritics.
        def draw(self):
            super().draw()
            self.canv.saveState()
            self.canv.translate(.22, 0)
            super().draw()
            self.canv.restoreState()
    def para(value, align=TA_LEFT, point=None, bold=False):
        style = ParagraphStyle('vc', fontName='VerconReport', fontSize=point or size,
                               leading=(point or size) * 1.22, alignment=align, spaceAfter=0)
        paragraph = EmphasizedParagraph if bold else Paragraph
        return paragraph(escape(str(value if value is not None else '')).replace('\n', '<br/>'), style)
    class PositionedBand(Flowable):
        def __init__(self, elements, available):
            super().__init__()
            self.width = available
            self.parts = []
            self.height = 0
            for element in elements:
                paragraph = para(element['text'], point=element.get('size', 8), bold=element.get('bold', False))
                _, height = paragraph.wrap(available * element['w'], 1000)
                self.parts.append((element, paragraph, height))
                self.height = max(self.height, element.get('y', 0) + max(element.get('h', 12), height))
        def draw(self):
            for element, paragraph, height in self.parts:
                paragraph.drawOn(self.canv, self.width * element['x'], self.height - element.get('y', 0) - height)
    def render(block, available):
        kind = block['t']
        if kind == 'positioned':
            return PositionedBand(block['elements'], available)
        if kind == 'break':
            return PageBreak()
        if kind == 'spacer':
            return Spacer(1, block.get('height', 5 * mm))
        if kind in ('title', 'text', 'group'):
            if kind == 'group':
                item = Table([[para(block['text'], TA_CENTER, size + .5, True)]], colWidths=[available])
                item.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), .4, colors.black), ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f7f8f2'))]))
                return item
            return para(block['text'], TA_CENTER if kind == 'title' else TA_LEFT,
                        layout['title_size'] if kind == 'title' else size, kind == 'title')
        if kind == 'box':
            data = [[para(label), para(value)] for label, value in block['rows']]
            item = Table(data, colWidths=[available * .34, available * .66])
        elif kind == 'sign':
            count = len(block['heads'])
            item = Table([[para(head, TA_CENTER) for head in block['heads']], [''] * count],
                         colWidths=[available / count] * count, rowHeights=[None, 14 * mm])
        elif kind == 'columns':
            count = len(block['children'])
            # Flatten parallel tables into one splittable table. Nesting complete
            # tables in a single outer row would fail on a long list of foods.
            children = [next(part for part in child if part['t'] == 'table') for child in block['children']]
            columns, groups, offset = [], [], 0
            for child, parts in zip(children, block['children']):
                total = sum(c['w'] for c in child['cols'])
                columns.extend([{**c, 'w': c['w'] / total} for c in child['cols']])
                label = next((p['text'] for p in parts if p['t'] == 'group'), '')
                groups.append({'label': label, 'start': offset, 'span': len(child['cols'])})
                offset += len(child['cols'])
            rows = []
            for index in range(max(len(child['rows']) for child in children)):
                cells = []
                for child in children:
                    cells.extend(child['rows'][index]['cells'] if index < len(child['rows']) else [''] * len(child['cols']))
                rows.append({'cells': cells})
            return render({'t': 'table', 'cols': columns, 'rows': rows,
                           'headers': [groups] if any(group['label'] for group in groups) else []}, available)
        elif kind == 'table':
            cols, commands, data = block['cols'], [], []
            headers = block.get('headers', [])
            for groups in headers:
                cells = [''] * len(cols)
                for group in groups:
                    start, span = group['start'], group['span']
                    cells[start] = para(group['label'], TA_CENTER)
                    if span > 1:
                        commands.append(('SPAN', (start, len(data)), (start + span - 1, len(data))))
                data.append(cells)
            data.append([para(c['label'], TA_CENTER, c.get('size'), True) for c in cols])
            head_count = len(data)
            for row in block['rows']:
                cells, index = [], 0
                for cell in row['cells']:
                    value, span = (cell['t'], cell.get('span', 1)) if isinstance(cell, dict) else (cell, 1)
                    col = cols[index]
                    cells.append(para(value, {'r': TA_RIGHT, 'c': TA_CENTER}.get(col.get('align'), TA_LEFT),
                                      col.get('size'), col.get('bold', False) or row.get('cls') == 'total'))
                    if span > 1:
                        commands.append(('SPAN', (index, len(data)), (index + span - 1, len(data))))
                        cells.extend([''] * (span - 1))
                    index += span
                data.append(cells)
                if row.get('cls') == 'total':
                    commands.append(('BACKGROUND', (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor('#ececec')))
            denominator = sum(c['w'] for c in cols)
            actual_width = available * block.get('width', 1)
            item = Table(data, colWidths=[actual_width * c['w'] / denominator for c in cols], repeatRows=head_count, hAlign='CENTER')
            commands.append(('BACKGROUND', (0, 0), (-1, head_count - 1), colors.HexColor('#ececec')))
            item.setStyle(TableStyle(commands))
        else:
            raise ValueError('Unknown report band: ' + kind)
        item.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .35, colors.black), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]))
        return item
    elements = identity['elements'] + layout.get('elements', [])
    values = {'UNITATE': identity['name'], 'CUI': identity['fiscal_code'], 'ADRESA': identity['address'],
              'ANTET': identity['header'], 'SUBSOL': identity['footer'], 'GESTIUNE': model['gestiune_name'],
              'PERIOADA': model['period_text'], 'RAPORT': model['title']}
    def band_height(band):
        ends = []
        for element in elements:
            if element['band'] != band:
                continue
            p = para(tokens(tokens(element['text'], {**values, 'PAGINA': 999, 'PAGINI': 999}), values),
                     point=element.get('size', 8), bold=element.get('bold', False))
            _, height = p.wrap(width * element['w'], 1000)
            ends.append(element.get('y', 0) + max(element.get('h', 12), height))
        return max(ends or [12])
    header_height, footer_height = band_height('header'), band_height('footer')
    class NumberedCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.states = []
        def showPage(self):
            self.states.append(dict(self.__dict__))
            self._startPage()
        def save(self):
            count = len(self.states)
            for state in self.states:
                self.__dict__.update(state)
                current = {**values, 'PAGINA': self._pageNumber, 'PAGINI': count}
                for e in elements:
                    if e['band'] not in ('header', 'footer'):
                        continue
                    y = page[1] - margin - float(e.get('y', 0)) if e['band'] == 'header' else margin + footer_height - float(e.get('y', 0))
                    p = para(tokens(tokens(e['text'], current), current), point=e.get('size', 8), bold=e.get('bold', False))
                    w, h = p.wrap(width * e['w'], 100 * mm)
                    p.drawOn(self, margin + width * e['x'], y - h)
                Canvas.showPage(self)
            Canvas.save(self)
    story = []
    for block in model['blocks']:
        story += [render(block, width), Spacer(1, 2 * mm)]
    stream = BytesIO()
    SimpleDocTemplate(stream, pagesize=page, leftMargin=margin, rightMargin=margin,
        topMargin=margin + header_height + 6, bottomMargin=margin + footer_height + 6,
        title=model['title']).build(story, canvasmaker=NumberedCanvas)
    return stream.getvalue()


def xlsx(model, column_widths=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, column_index_from_string
    from . import report_excel
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Raport'
    row_index = 1
    widths = ColumnWidths(column_widths)
    edge = Side(style='thin', color='AAAAAA')
    def write(row, column, value, bold=False, point=10):
        cell = sheet.cell(row, column, str(value if value is not None else ''))
        # Exact formatted decimals and identifiers stay literal, never formulas
        # or binary-float approximations imposed by Excel's 15 digit limit.
        cell.data_type = 's'
        cell.font = Font(name='Arial', size=point, bold=bold)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        return cell
    def blocks(items):
        nonlocal row_index
        for block in items:
            kind = block['t']
            if kind == 'positioned':
                for element in sorted(block['elements'], key=lambda e: (e.get('y', 0), e.get('x', 0))):
                    write(row_index, 1, element['text'], element.get('bold', False), element.get('size', 10))
                    row_index += 1
            elif kind == 'columns':
                for child in block['children']:
                    blocks(child)
            elif kind in ('title', 'text', 'group'):
                write(row_index, 1, block['text'], kind != 'text')
                row_index += 2
            elif kind in ('box', 'sign'):
                for cells in block.get('rows', [block.get('heads', [])]):
                    for i, value in enumerate(cells, 1):
                        write(row_index, i, value)
                    row_index += 1
                row_index += 1
            elif kind == 'table':
                for i, col in enumerate(block['cols'], 1):
                    write(row_index, i, col['label'], True, col.get('size', 10))
                    dimension = sheet.column_dimensions[get_column_letter(i)]
                    dimension.width = widths.observe(i, report_excel.key(col, i - 1, block['cols']), col['w'], dimension.width)
                row_index += 1
                for record in block['rows']:
                    for i, value in enumerate(record['cells'], 1):
                        col = block['cols'][i - 1]
                        cell = write(row_index, i, value['t'] if isinstance(value, dict) else value,
                                     record.get('cls') == 'total' or col.get('bold', False), col.get('size', 10))
                        cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                        if record.get('cls') == 'total':
                            cell.fill = PatternFill('solid', fgColor='ECECEC')
                    row_index += 1
                row_index += 2
    identity = model['identity']
    values = {'UNITATE': identity['name'], 'CUI': identity['fiscal_code'], 'ADRESA': identity['address'],
              'GESTIUNE': model['gestiune_name'], 'PERIOADA': model['period_text'], 'RAPORT': model['title'],
              'ANTET': identity['header'], 'SUBSOL': identity['footer']}
    def page_band(name):
        nonlocal row_index
        elements = [e for e in identity['elements'] + model['layout'].get('elements', []) if e['band'] == name]
        for element in sorted(elements, key=lambda e: (e.get('y', 0), e.get('x', 0))):
            text = tokens(tokens(element['text'], values), values)
            if '{PAGINA}' in text or '{PAGINI}' in text:
                part = sheet.oddHeader if name == 'header' else sheet.oddFooter
                position = part.right if element.get('x', 0) > .65 else part.left
                position.text = text.replace('{PAGINA}', '&P').replace('{PAGINI}', '&N')
            else:
                write(row_index, 1, text, element.get('bold', name == 'header'), element.get('size', 10))
                row_index += 1
        row_index += 1
    page_band('header')
    blocks(model['blocks'])
    page_band('footer')
    for dimension in sheet.column_dimensions.values():
        index = column_index_from_string(dimension.index)
        dimension.width = widths.final(index, dimension.width)
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = 'landscape' if model['layout']['landscape'] else 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.freeze_panes = 'A4'
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
